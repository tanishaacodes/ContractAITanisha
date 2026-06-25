"""
Document Generator for Contract Redlining
==========================================
Generates:
- DOCX files with Word-compatible track changes styling
- PDF files with redline annotations
"""

import os
import tempfile
from datetime import datetime
from typing import List, Dict, Any, Optional
from io import BytesIO

from django.conf import settings

# DOCX generation
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# PDF generation
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.pdfgen import canvas


def generate_docx_redline(
    clauses: List[Dict[str, Any]],
    contract_name: str = "Contract",
    output_path: Optional[str] = None
) -> str:
    """
    Generate a Word document with redline track changes styling.

    Each clause shows:
    - Original text (red, strikethrough)
    - Suggested text (green, underlined)
    - Legal reasoning (italic, boxed)

    Returns: Path to generated DOCX file
    """
    doc = Document()

    # Set document properties
    core_props = doc.core_properties
    core_props.title = f"{contract_name} - AI Redlined Version"
    core_props.author = "ContractAI Redline Engine"
    core_props.created = datetime.now()

    # Add title
    title = doc.add_heading(f'{contract_name}', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run('AI-Generated Redline Review')
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor(100, 100, 100)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add generation timestamp
    timestamp = doc.add_paragraph()
    timestamp_run = timestamp.add_run(f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    timestamp_run.font.size = Pt(10)
    timestamp_run.font.color.rgb = RGBColor(128, 128, 128)
    timestamp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()  # Spacer

    # Add summary section
    doc.add_heading('Executive Summary', level=1)

    # Calculate statistics
    total_clauses = len(clauses)
    high_risk = sum(1 for c in clauses if c.get('risk_score', 0) >= 70)
    medium_risk = sum(1 for c in clauses if 40 <= c.get('risk_score', 0) < 70)
    low_risk = sum(1 for c in clauses if c.get('risk_score', 0) < 40)
    accepted = sum(1 for c in clauses if c.get('status') == 'ACCEPTED')

    summary_table_data = [
        ['Total Clauses Reviewed', str(total_clauses)],
        ['High Risk Clauses', str(high_risk)],
        ['Medium Risk Clauses', str(medium_risk)],
        ['Low Risk Clauses', str(low_risk)],
        ['Changes Accepted', str(accepted)],
    ]

    table = doc.add_table(rows=len(summary_table_data), cols=2)
    table.style = 'Table Grid'
    for i, row_data in enumerate(summary_table_data):
        row = table.rows[i]
        row.cells[0].text = row_data[0]
        row.cells[1].text = row_data[1]

    doc.add_paragraph()  # Spacer
    doc.add_page_break()

    # Add each clause with redline formatting
    doc.add_heading('Detailed Clause Review', level=1)

    for i, clause in enumerate(clauses, 1):
        # Clause header
        clause_heading = doc.add_heading(f'{i}. {clause.get("clause_name", f"Clause {i}")}', level=2)

        # Risk badge
        risk_score = clause.get('risk_score', 0)
        risk_type = clause.get('risk_type', 'OTHER')

        if risk_score >= 70:
            risk_color = RGBColor(220, 53, 69)  # Red
            risk_level = 'HIGH RISK'
        elif risk_score >= 40:
            risk_color = RGBColor(255, 193, 7)  # Yellow/Orange
            risk_level = 'MEDIUM RISK'
        else:
            risk_color = RGBColor(40, 167, 69)  # Green
            risk_level = 'LOW RISK'

        risk_para = doc.add_paragraph()
        risk_run = risk_para.add_run(f'{risk_level} ({risk_score}/100) - {risk_type}')
        risk_run.font.bold = True
        risk_run.font.color.rgb = risk_color

        # Original text (strikethrough, red)
        doc.add_paragraph().add_run('ORIGINAL:').bold = True

        original_para = doc.add_paragraph()
        original_run = original_para.add_run(clause.get('original_text', ''))
        original_run.font.strike = True
        original_run.font.color.rgb = RGBColor(255, 0, 0)

        # Suggested text (green)
        doc.add_paragraph().add_run('SUGGESTED:').bold = True

        suggested_para = doc.add_paragraph()
        suggested_run = suggested_para.add_run(clause.get('suggested_text', clause.get('original_text', '')))
        suggested_run.font.color.rgb = RGBColor(0, 128, 0)
        suggested_run.font.underline = True

        # Add accepted text if different
        accepted_text = clause.get('accepted_text')
        if accepted_text and accepted_text != clause.get('suggested_text'):
            doc.add_paragraph().add_run('ACCEPTED (EDITED):').bold = True
            accepted_para = doc.add_paragraph()
            accepted_run = accepted_para.add_run(accepted_text)
            accepted_run.font.color.rgb = RGBColor(0, 100, 0)
            accepted_run.font.bold = True

        # Legal reasoning box
        if clause.get('court_reasoning') or clause.get('risk_explanation'):
            legal_para = doc.add_paragraph()
            legal_para.paragraph_format.left_indent = Inches(0.5)

            # Add border effect using background shading
            shading_elm = OxmlElement('w:shd')
            shading_elm.set(qn('w:fill'), 'F5F5F5')
            legal_para._p.get_or_add_pPr().append(shading_elm)

            legal_header = legal_para.add_run('Legal Analysis: ')
            legal_header.font.bold = True
            legal_header.font.size = Pt(10)

            legal_text = clause.get('court_reasoning') or clause.get('risk_explanation', '')
            legal_content = legal_para.add_run(legal_text)
            legal_content.font.italic = True
            legal_content.font.size = Pt(10)
            legal_content.font.color.rgb = RGBColor(64, 64, 64)

        # Status indicator
        status = clause.get('status', 'PENDING')
        status_para = doc.add_paragraph()
        status_text = {'ACCEPTED': 'Accepted', 'REJECTED': 'Rejected', 'MODIFIED': 'Modified', 'PENDING': 'Pending Review'}
        status_run = status_para.add_run(f'Status: {status_text.get(status, status)}')
        status_run.font.size = Pt(9)
        status_run.font.italic = True

        doc.add_paragraph()  # Spacer between clauses

    # Generate output path
    if not output_path:
        output_dir = os.path.join(settings.MEDIA_ROOT, 'redlines')
        os.makedirs(output_dir, exist_ok=True)
        timestamp_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        safe_name = "".join(c for c in contract_name if c.isalnum() or c in (' ', '-', '_')).rstrip()
        output_path = os.path.join(output_dir, f'{safe_name}_Redlined_{timestamp_str}.docx')

    doc.save(output_path)
    return output_path


def generate_pdf_redline(
    clauses: List[Dict[str, Any]],
    contract_name: str = "Contract",
    output_path: Optional[str] = None
) -> str:
    """
    Generate a PDF document with redline formatting.

    Uses visual indicators:
    - Red text with strikethrough for deleted content
    - Green text for added content
    - Gray boxes for legal reasoning

    Returns: Path to generated PDF file
    """
    # Generate output path
    if not output_path:
        output_dir = os.path.join(settings.MEDIA_ROOT, 'redlines')
        os.makedirs(output_dir, exist_ok=True)
        timestamp_str = datetime.now().strftime('%Y%m%d_%H%M%S')
        safe_name = "".join(c for c in contract_name if c.isalnum() or c in (' ', '-', '_')).rstrip()
        output_path = os.path.join(output_dir, f'{safe_name}_Redlined_{timestamp_str}.pdf')

    # Create PDF document
    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=72,
        leftMargin=72,
        topMargin=72,
        bottomMargin=72
    )

    # Define styles
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        spaceAfter=12,
        alignment=1  # Center
    )

    subtitle_style = ParagraphStyle(
        'CustomSubtitle',
        parent=styles['Normal'],
        fontSize=12,
        textColor=colors.gray,
        spaceAfter=20,
        alignment=1
    )

    heading_style = ParagraphStyle(
        'ClauseHeading',
        parent=styles['Heading2'],
        fontSize=14,
        spaceAfter=6,
        textColor=colors.HexColor('#333333')
    )

    deleted_style = ParagraphStyle(
        'Deleted',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.red,
        spaceAfter=8
    )

    added_style = ParagraphStyle(
        'Added',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#006400'),
        spaceAfter=8
    )

    legal_style = ParagraphStyle(
        'Legal',
        parent=styles['Normal'],
        fontSize=9,
        textColor=colors.HexColor('#404040'),
        backColor=colors.HexColor('#F5F5F5'),
        leftIndent=20,
        rightIndent=20,
        spaceAfter=12,
        spaceBefore=6
    )

    risk_high_style = ParagraphStyle(
        'RiskHigh',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#DC3545'),
        spaceBefore=4,
        spaceAfter=8
    )

    risk_medium_style = ParagraphStyle(
        'RiskMedium',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#FFC107'),
        spaceBefore=4,
        spaceAfter=8
    )

    risk_low_style = ParagraphStyle(
        'RiskLow',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#28A745'),
        spaceBefore=4,
        spaceAfter=8
    )

    # Build document content
    story = []

    # Title
    story.append(Paragraph(f'{contract_name}', title_style))
    story.append(Paragraph('AI-Generated Redline Review', subtitle_style))
    story.append(Paragraph(f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}', subtitle_style))
    story.append(Spacer(1, 20))

    # Summary statistics
    total_clauses = len(clauses)
    high_risk = sum(1 for c in clauses if c.get('risk_score', 0) >= 70)
    medium_risk = sum(1 for c in clauses if 40 <= c.get('risk_score', 0) < 70)
    low_risk = sum(1 for c in clauses if c.get('risk_score', 0) < 40)

    story.append(Paragraph('Executive Summary', styles['Heading2']))

    summary_data = [
        ['Metric', 'Count'],
        ['Total Clauses', str(total_clauses)],
        ['High Risk', str(high_risk)],
        ['Medium Risk', str(medium_risk)],
        ['Low Risk', str(low_risk)],
    ]

    summary_table = Table(summary_data, colWidths=[3*inch, 1.5*inch])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#333333')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F8F9FA')),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#DEE2E6')),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('TOPPADDING', (0, 1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 8),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 30))
    story.append(PageBreak())

    # Detailed clause review
    story.append(Paragraph('Detailed Clause Review', styles['Heading1']))
    story.append(Spacer(1, 12))

    for i, clause in enumerate(clauses, 1):
        # Clause header
        clause_name = clause.get('clause_name', f'Clause {i}')
        story.append(Paragraph(f'{i}. {clause_name}', heading_style))

        # Risk indicator
        risk_score = clause.get('risk_score', 0)
        risk_type = clause.get('risk_type', 'OTHER')

        if risk_score >= 70:
            risk_style = risk_high_style
            risk_label = 'HIGH RISK'
        elif risk_score >= 40:
            risk_style = risk_medium_style
            risk_label = 'MEDIUM RISK'
        else:
            risk_style = risk_low_style
            risk_label = 'LOW RISK'

        story.append(Paragraph(f'<b>{risk_label}</b> ({risk_score}/100) - {risk_type}', risk_style))

        # Original text (marked as deleted)
        original_text = clause.get('original_text', '').replace('\n', '<br/>')
        story.append(Paragraph(f'<b>DELETED:</b> <strike>{original_text}</strike>', deleted_style))

        # Suggested text
        suggested_text = clause.get('suggested_text', original_text).replace('\n', '<br/>')
        story.append(Paragraph(f'<b>ADDED:</b> <u>{suggested_text}</u>', added_style))

        # Legal reasoning
        legal_text = clause.get('court_reasoning') or clause.get('risk_explanation', '')
        if legal_text:
            legal_text = legal_text.replace('\n', '<br/>')
            story.append(Paragraph(f'<b>Legal Analysis:</b> <i>{legal_text}</i>', legal_style))

        story.append(Spacer(1, 20))

    # Build PDF
    doc.build(story)
    return output_path


def generate_track_changes_summary(clauses: List[Dict[str, Any]]) -> Dict[str, Any]:
    """
    Generate a summary of all track changes for export.
    """
    total = len(clauses)
    accepted = sum(1 for c in clauses if c.get('status') == 'ACCEPTED')
    rejected = sum(1 for c in clauses if c.get('status') == 'REJECTED')
    modified = sum(1 for c in clauses if c.get('status') == 'MODIFIED')
    pending = sum(1 for c in clauses if c.get('status') == 'PENDING')

    risk_breakdown = {
        'high': sum(1 for c in clauses if c.get('risk_score', 0) >= 70),
        'medium': sum(1 for c in clauses if 40 <= c.get('risk_score', 0) < 70),
        'low': sum(1 for c in clauses if c.get('risk_score', 0) < 40)
    }

    return {
        'total_clauses': total,
        'accepted': accepted,
        'rejected': rejected,
        'modified': modified,
        'pending': pending,
        'completion_rate': round((accepted + rejected + modified) / total * 100, 1) if total > 0 else 0,
        'acceptance_rate': round(accepted / total * 100, 1) if total > 0 else 0,
        'risk_breakdown': risk_breakdown,
        'average_risk_score': round(sum(c.get('risk_score', 0) for c in clauses) / total, 1) if total > 0 else 0
    }
