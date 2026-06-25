"""
Clause Edit Service
Handles clause-level editing with version tracking and risk re-analysis
"""

import logging
from typing import Dict, List, Optional
from datetime import datetime
from django.db import transaction
from django.db.models import Max

from core.models import Contract, Clause, ClauseVersion, ContractVersion, User
from .risk_scoring_model import ClauseRiskEnricher

logger = logging.getLogger(__name__)

# Initialize risk enricher
clause_risk_enricher = ClauseRiskEnricher()


class ClauseEditService:
    """Service for managing clause edits and contract regeneration"""

    def get_contract_clauses_for_editing(self, contract_id: str) -> Dict:
        """
        Get all clauses from a contract for editing

        Args:
            contract_id: UUID of the contract

        Returns:
            Dict with contract info and clauses list
        """
        try:
            contract = Contract.objects.get(id=contract_id)
            clauses = Clause.objects.filter(contract=contract, found=True).order_by('clause_name')

            clause_list = []
            for clause in clauses:
                # Get latest version if exists
                latest_version = ClauseVersion.objects.filter(clause=clause).order_by('-version_number').first()

                clause_data = {
                    'id': clause.id,
                    'clause_name': clause.clause_name,
                    'clause_type': clause.clause_type or 'Unknown',
                    'original_text': clause.extracted_text,
                    'current_text': latest_version.modified_text if latest_version else clause.extracted_text,
                    'risk_score': clause.risk_score,
                    'risk_level': clause.risk_level,
                    'likelihood_score': clause.likelihood_score,
                    'impact_score': clause.impact_score,
                    'version_count': ClauseVersion.objects.filter(clause=clause).count(),
                    'last_modified': latest_version.modified_at if latest_version else None,
                    'has_been_edited': latest_version is not None,
                }
                clause_list.append(clause_data)

            return {
                'success': True,
                'contract': {
                    'id': contract.id,
                    'filename': contract.original_filename,
                    'status': contract.status,
                    'can_edit': contract.status == 'DRAFT',  # Only draft contracts can be edited
                },
                'clauses': clause_list,
                'total_clauses': len(clause_list),
            }

        except Contract.DoesNotExist:
            return {
                'success': False,
                'error': 'Contract not found'
            }
        except Exception as e:
            logger.error(f"Error getting clauses for editing: {str(e)}")
            return {
                'success': False,
                'error': str(e)
            }

    @transaction.atomic
    def update_clause(
        self,
        clause_id: str,
        modified_text: str,
        change_description: str,
        user_id: str
    ) -> Dict:
        """
        Update a clause with new text, create version history, and re-run risk analysis

        Args:
            clause_id: UUID of the clause
            modified_text: New clause text
            change_description: User's explanation of the change
            user_id: User making the change

        Returns:
            Dict with success status and updated clause data
        """
        try:
            clause = Clause.objects.get(id=clause_id)
            user = User.objects.get(id=user_id)

            # Get current text (either from latest version or original)
            latest_version = ClauseVersion.objects.filter(clause=clause).order_by('-version_number').first()
            original_text = latest_version.modified_text if latest_version else clause.extracted_text

            # Check if text actually changed
            if modified_text.strip() == original_text.strip():
                return {
                    'success': False,
                    'error': 'No changes detected in clause text'
                }

            # Get next version number
            max_version = ClauseVersion.objects.filter(clause=clause).aggregate(Max('version_number'))
            next_version = (max_version['version_number__max'] or 0) + 1

            # Re-run risk analysis on modified clause
            risk_analysis = self._analyze_modified_clause(clause, modified_text)

            # Create new clause version
            clause_version = ClauseVersion.objects.create(
                clause=clause,
                version_number=next_version,
                original_text=original_text,
                modified_text=modified_text,
                change_description=change_description or f"Modified clause (version {next_version})",
                new_risk_score=risk_analysis['risk_score'],
                new_risk_level=risk_analysis['risk_level'],
                new_likelihood_score=risk_analysis['likelihood_score'],
                new_impact_score=risk_analysis['impact_score'],
                modified_by=user,
            )

            # Update the clause with new risk scores (optional - keep original or update)
            # For now, we'll keep original scores and only store new scores in version
            # If you want to update the clause itself, uncomment below:
            # clause.extracted_text = modified_text
            # clause.risk_score = risk_analysis['risk_score']
            # clause.risk_level = risk_analysis['risk_level']
            # clause.likelihood_score = risk_analysis['likelihood_score']
            # clause.impact_score = risk_analysis['impact_score']
            # clause.save()

            logger.info(f"Clause {clause_id} updated to version {next_version} by user {user_id}")

            return {
                'success': True,
                'clause_version': {
                    'id': clause_version.id,
                    'version_number': clause_version.version_number,
                    'modified_text': clause_version.modified_text,
                    'change_description': clause_version.change_description,
                    'original_risk_score': clause.risk_score,
                    'new_risk_score': clause_version.new_risk_score,
                    'original_risk_level': clause.risk_level,
                    'new_risk_level': clause_version.new_risk_level,
                    'risk_improved': (clause.risk_score or 1.0) > (clause_version.new_risk_score or 0),
                    'modified_at': clause_version.modified_at,
                    'modified_by': user.email,
                },
                'clause': {
                    'id': clause.id,
                    'clause_name': clause.clause_name,
                    'version_count': next_version,
                }
            }

        except Clause.DoesNotExist:
            return {'success': False, 'error': 'Clause not found'}
        except User.DoesNotExist:
            return {'success': False, 'error': 'User not found'}
        except Exception as e:
            logger.error(f"Error updating clause: {str(e)}")
            return {'success': False, 'error': str(e)}

    def _analyze_modified_clause(self, clause: Clause, modified_text: str) -> Dict:
        """
        Run risk analysis on modified clause text

        Args:
            clause: Original Clause object
            modified_text: Modified clause text

        Returns:
            Dict with risk scores
        """
        try:
            # Use ClauseRiskEnricher to score the modified clause
            enrichment = clause_risk_enricher.enrich_risk(
                clause_text=modified_text,
                intent_name=clause.clause_type or 'General',
                party="COUNTERPARTY",
                base_risk=0.5
            )

            # Calculate likelihood and impact scores (1-5 scale)
            clause_strength = enrichment.get('clause_strength', 0)
            party_bias = enrichment.get('party_bias', 0)
            missing_safeguards = enrichment.get('missing_safeguards', 0)
            financial_factor = enrichment.get('financial_factor', 1.0)
            final_risk = enrichment.get('final_risk', 0.5)

            # Likelihood: combination of clause strength and party bias
            combined_likelihood = (clause_strength + party_bias) / 2
            likelihood = min(max(1.0, combined_likelihood * 3.0), 5.0)

            # Impact: combination of missing safeguards and financial exposure
            combined_impact = (missing_safeguards + (financial_factor - 1.0)) / 2
            impact = min(max(1.0, combined_impact * final_risk * 5.0), 5.0)

            # Risk level classification
            if final_risk < 0.4:
                risk_level = 'LOW'
            elif final_risk < 0.7:
                risk_level = 'MEDIUM'
            else:
                risk_level = 'HIGH'

            return {
                'risk_score': round(final_risk, 3),
                'risk_level': risk_level,
                'likelihood_score': round(likelihood, 2),
                'impact_score': round(impact, 2),
                'risk_factors': enrichment,
            }

        except Exception as e:
            logger.error(f"Error analyzing modified clause: {str(e)}")
            # Return neutral scores on error
            return {
                'risk_score': 0.5,
                'risk_level': 'MEDIUM',
                'likelihood_score': 2.5,
                'impact_score': 2.5,
                'risk_factors': {},
            }

    def get_clause_version_history(self, clause_id: str) -> Dict:
        """
        Get version history for a clause

        Args:
            clause_id: UUID of the clause

        Returns:
            Dict with version history
        """
        try:
            clause = Clause.objects.get(id=clause_id)
            versions = ClauseVersion.objects.filter(clause=clause).order_by('-version_number')

            version_list = []
            for version in versions:
                version_list.append({
                    'id': version.id,
                    'version_number': version.version_number,
                    'original_text': version.original_text,
                    'modified_text': version.modified_text,
                    'change_description': version.change_description,
                    'risk_score': version.new_risk_score,
                    'risk_level': version.new_risk_level,
                    'modified_by': version.modified_by.email,
                    'modified_at': version.modified_at,
                    'approved_by_id': version.approved_by_id,
                    'approved_at': version.approved_at,
                })

            return {
                'success': True,
                'clause': {
                    'id': clause.id,
                    'clause_name': clause.clause_name,
                    'original_text': clause.extracted_text,
                },
                'versions': version_list,
                'total_versions': len(version_list),
            }

        except Clause.DoesNotExist:
            return {'success': False, 'error': 'Clause not found'}
        except Exception as e:
            logger.error(f"Error getting version history: {str(e)}")
            return {'success': False, 'error': str(e)}


    def _generate_docx_contract(self, contract, clauses, user, version_number, uploads_dir, base_name, modified_clause_count):
        """
        Generate a DOCX version of the contract

        Returns:
            Tuple of (file_path, filename, full_modified_text)
        """
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import os

        # Create document
        doc = Document()

        # Title
        title = doc.add_heading(f"{contract.original_filename} - Version {version_number}", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.runs[0]
        title_run.font.color.rgb = RGBColor(0, 0, 139)  # Dark blue

        # Metadata section
        doc.add_paragraph()
        metadata = doc.add_paragraph()
        metadata_run = metadata.add_run(f"Modified by: {user.email}\n")
        metadata_run.font.size = Pt(10)
        metadata_run.font.color.rgb = RGBColor(128, 128, 128)

        metadata_run2 = metadata.add_run(f"Modified clauses: {modified_clause_count}")
        metadata_run2.font.size = Pt(10)
        metadata_run2.font.color.rgb = RGBColor(128, 128, 128)

        doc.add_paragraph()
        doc.add_paragraph("_" * 80)
        doc.add_paragraph()

        # Contract content
        contract_parts = []
        for clause in clauses:
            latest_version = ClauseVersion.objects.filter(clause=clause).order_by('-version_number').first()

            # Add clause heading
            heading = doc.add_heading(clause.clause_name, level=2)

            # Add clause text
            if latest_version:
                clause_text = latest_version.modified_text
                contract_parts.append(f"\n{clause.clause_name}\n{'='*50}\n{clause_text}\n")
            else:
                clause_text = clause.extracted_text
                contract_parts.append(f"\n{clause.clause_name}\n{'='*50}\n{clause_text}\n")

            # Add text paragraph
            para = doc.add_paragraph(clause_text)
            para.paragraph_format.space_after = Pt(12)
            para.paragraph_format.line_spacing = 1.5

            # Add separator
            doc.add_paragraph()

        full_modified_text = "\n".join(contract_parts)

        # Save DOCX
        new_filename = f"{base_name}_v{version_number}.docx"
        file_path = os.path.join(uploads_dir, new_filename)
        doc.save(file_path)

        return file_path, new_filename, full_modified_text

    @transaction.atomic
    def regenerate_contract(self, contract_id: str, user_id: str, output_format: str = 'pdf') -> Dict:
        """
        Regenerate contract with all modified clauses

        Args:
            contract_id: UUID of the contract
            user_id: User requesting regeneration
            output_format: 'pdf' or 'docx' (default: 'pdf')

        Returns:
            Dict with new contract version info and download path
        """
        try:
            import os
            from django.conf import settings
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch

            contract = Contract.objects.get(id=contract_id)
            user = User.objects.get(id=user_id)

            # Get all clauses with their latest versions
            clauses = Clause.objects.filter(contract=contract, found=True).order_by('id')

            # Check if ANY clauses have been modified
            modified_versions = ClauseVersion.objects.filter(
                clause__contract=contract,
                clause__found=True
            ).exists()

            if not modified_versions:
                return {
                    'success': False,
                    'error': 'No modified clauses found to regenerate contract'
                }

            # Build contract with modified clauses
            contract_parts = []
            change_summary_parts = []
            modified_clause_count = 0

            for clause in clauses:
                latest_version = ClauseVersion.objects.filter(clause=clause).order_by('-version_number').first()

                if latest_version:
                    # Use modified text
                    contract_parts.append(f"\n{clause.clause_name}\n{'='*50}\n{latest_version.modified_text}\n")
                    modified_clause_count += 1
                    change_summary_parts.append(
                        f"{clause.clause_name}: {latest_version.change_description}"
                    )
                else:
                    # Use original text
                    contract_parts.append(f"\n{clause.clause_name}\n{'='*50}\n{clause.extracted_text}\n")

            full_modified_text = "\n".join(contract_parts)

            # Create new version number
            max_version = ContractVersion.objects.filter(contract=contract).aggregate(Max('version_number'))
            next_version = (max_version['version_number__max'] or 0) + 1

            # Generate filename
            base_name = os.path.splitext(contract.original_filename)[0]

            # Create uploads directory if it doesn't exist
            uploads_dir = os.path.join(settings.BASE_DIR, 'uploads')
            os.makedirs(uploads_dir, exist_ok=True)

            # Generate contract based on format
            if output_format.lower() == 'docx':
                # Generate DOCX
                file_path, new_filename, full_modified_text = self._generate_docx_contract(
                    contract, clauses, user, next_version, uploads_dir, base_name, modified_clause_count
                )
                file_type = 'docx'
                content_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            else:
                # Generate PDF (default)
                new_filename = f"{base_name}_v{next_version}.pdf"
                file_path = os.path.join(uploads_dir, new_filename)
                file_type = 'pdf'
                content_type = 'application/pdf'

                # Build contract with modified clauses
                contract_parts = []
                for clause in clauses:
                    latest_version = ClauseVersion.objects.filter(clause=clause).order_by('-version_number').first()
                    if latest_version:
                        contract_parts.append(f"\n{clause.clause_name}\n{'='*50}\n{latest_version.modified_text}\n")
                    else:
                        contract_parts.append(f"\n{clause.clause_name}\n{'='*50}\n{clause.extracted_text}\n")

                full_modified_text = "\n".join(contract_parts)

                # Generate PDF with modified contract
                doc = SimpleDocTemplate(file_path, pagesize=letter,
                                        rightMargin=72, leftMargin=72,
                                        topMargin=72, bottomMargin=18)

                styles = getSampleStyleSheet()
                story = []

                # Title
                title_style = ParagraphStyle(
                    'CustomTitle',
                    parent=styles['Heading1'],
                    fontSize=24,
                    textColor='darkblue',
                    spaceAfter=30,
                    alignment=1  # Center
                )
                story.append(Paragraph(f"{contract.original_filename} - Version {next_version}", title_style))
                story.append(Spacer(1, 0.2 * inch))

                # Version info
                info_style = ParagraphStyle(
                    'VersionInfo',
                    parent=styles['Normal'],
                    fontSize=10,
                    textColor='gray'
                )
                story.append(Paragraph(f"Modified by: {user.email}", info_style))
                story.append(Paragraph(f"Modified clauses: {modified_clause_count}", info_style))
                story.append(Spacer(1, 0.3 * inch))

                # Contract content
                content_style = ParagraphStyle(
                    'Content',
                    parent=styles['Normal'],
                    fontSize=11,
                    leading=14,
                    spaceAfter=10
                )

                # Split text into paragraphs and add to story
                paragraphs = full_modified_text.split('\n\n')
                for para in paragraphs:
                    if para.strip():
                        story.append(Paragraph(para.strip(), content_style))
                        story.append(Spacer(1, 0.1 * inch))

                # Build PDF
                doc.build(story)

            # Create ContractVersion record
            contract_version = ContractVersion.objects.create(
                contract=contract,
                version_number=next_version,
                created_by=user,
                change_description=f"Modified {modified_clause_count} clauses",
                filename=new_filename,
                original_filename=new_filename,
                file_type=file_type,
                file_path=file_path,
                full_text=full_modified_text,
                contract_type=contract.contract_type,
                contract_value=contract.contract_value,
                party_name=contract.party_name,
                contract_duration=contract.contract_duration,
            )

            logger.info(f"Contract {contract_id} regenerated as version {next_version}")

            return {
                'success': True,
                'contract_version': {
                    'id': contract_version.id,
                    'version_number': contract_version.version_number,
                    'filename': contract_version.filename,
                    'file_path': file_path,
                    'modified_clause_count': modified_clause_count,
                    'file_size': os.path.getsize(file_path),
                    'created_at': contract_version.created_at,
                    'created_by': user.email,
                    'change_summary': '\n'.join(change_summary_parts[:10]),  # First 10 changes
                }
            }

        except Contract.DoesNotExist:
            return {'success': False, 'error': 'Contract not found'}
        except User.DoesNotExist:
            return {'success': False, 'error': 'User not found'}
        except Exception as e:
            logger.error(f"Error regenerating contract: {str(e)}")
            import traceback
            logger.error(traceback.format_exc())
            return {'success': False, 'error': str(e)}


# Singleton instance
clause_edit_service = ClauseEditService()
